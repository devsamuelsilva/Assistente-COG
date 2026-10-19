from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

import numpy as np

from app.buscar import buscar, expandir
from app.chunking import Trecho, conferir_sumario, fatiar, linha_documento
from app.extrair import Documento, Figura, _compactar_layout, anexar_figuras, carregar_acervo, carregar_arquivo, figuras_de_conteudo
from app.indice import Indice, _com_cache
from app.visao import transcrever_figuras, transcricoes_guardadas
from app.buscar import _com_continuacao

RAIZ = Path(__file__).resolve().parents[1]
EXEMPLO = RAIZ / "tests" / "dados" / "exemplo-normativo.md"


def _documento(paginas: list[tuple[int, str]], codigo: str = "ITOM-01", titulo: str = "ITOM") -> Documento:
    return Documento(
        arquivo="itom.pdf",
        caminho=Path("itom.pdf"),
        texto="\n".join(texto for _, texto in paginas),
        paginas=paginas,
        codigo=codigo,
        revisao="00",
        vigente=True,
        titulo=titulo,
    )


class ExtrairChunkingTeste(unittest.TestCase):
    def test_exemplo_tem_secoes_e_valor(self) -> None:
        documento = carregar_arquivo(EXEMPLO)
        self.assertEqual(documento.codigo, "IT-EXEMPLO-001")
        self.assertEqual(documento.revisao, "0")
        self.assertTrue(documento.vigente)
        self.assertTrue(documento.exemplo)
        trechos = fatiar(documento)
        secoes = {trecho.secao for trecho in trechos}
        self.assertIn("2. Valor de teste", secoes)
        valor = next(trecho for trecho in trechos if trecho.secao == "2. Valor de teste")
        self.assertIn("42 unidades de exemplo", valor.texto)

    def test_so_o_exemplo_fica_quando_nao_ha_normativo_real(self) -> None:
        with tempfile.TemporaryDirectory() as temporario:
            pasta = Path(temporario)
            (pasta / EXEMPLO.name).write_bytes(EXEMPLO.read_bytes())
            documentos = carregar_acervo(pasta)
        self.assertEqual(len(documentos), 1)
        self.assertTrue(documentos[0].exemplo)

    def test_normativo_real_deixa_o_exemplo_de_fora(self) -> None:
        with tempfile.TemporaryDirectory() as temporario:
            pasta = Path(temporario)
            (pasta / "exemplo-normativo.md").write_text(
                "# Exemplo\n\nCódigo: IT-EXEMPLO-001\nRevisão: 0\nStatus: vigente\n\n## 1. Teste\n\nTexto.\n",
                encoding="utf-8",
            )
            (pasta / "it-real.md").write_text(
                "# Instrução real\n\nCódigo: IT-100\nRevisão: 2\nStatus: substituído\n\n## 1. Escopo\n\nTexto real.\n",
                encoding="utf-8",
            )
            (pasta / "it-vigente.md").write_text(
                "# Instrução vigente\n\nCódigo: IT-200\nRevisão: 1\nStatus: vigente\n\n## 4. Manobra de teste\n\nPasso escrito.\n",
                encoding="utf-8",
            )
            documentos = carregar_acervo(pasta)
            codigos = {documento.codigo for documento in documentos}
            self.assertEqual(codigos, {"IT-100", "IT-200"})
            vigente = next(documento for documento in documentos if documento.codigo == "IT-200")
            substituido = next(documento for documento in documentos if documento.codigo == "IT-100")
            self.assertTrue(vigente.vigente)
            self.assertFalse(substituido.vigente)
            self.assertIn("4. Manobra de teste", {trecho.secao for trecho in fatiar(vigente)})

    def test_item_numerado_fica_inteiro_e_sem_cabecalho_repetido(self) -> None:
        cabecalho = "Normativo Operacional – Pós-Operação"
        pagina_1 = "\n".join(
            [
                cabecalho,
                "2.3. A Pós-Operação elabora o RAO quando",
                "a) desligamento acidental",
                "i) desligamentos não programados com duração superior a 12 horas",
            ]
        )
        pagina_2 = "\n".join(
            [
                cabecalho,
                "2.4. A Pós-Operação elabora o RAI quando",
                "a) falha humana durante a intervenção",
                "b) desempenho insatisfatório da equipe",
                "c) recursos de suporte à operação",
                "d) desligamento programado com duração superior a 12 horas",
                "e) manutenções gerais nas subestações",
            ]
        )
        documento = Documento(
            arquivo="no-01.pdf",
            caminho=Path("no-01.pdf"),
            texto="",
            paginas=[(1, pagina_1), (2, pagina_2)],
            codigo="NO-01",
            revisao="00",
            vigente=True,
            titulo="Pós-Operação",
        )
        trechos = fatiar(documento)
        rai = next(trecho for trecho in trechos if trecho.secao.startswith("2.4"))
        self.assertIn("falha humana durante a intervenção", rai.texto)
        self.assertIn("manutenções gerais nas subestações", rai.texto)
        self.assertNotIn(cabecalho, rai.texto)
        self.assertNotIn("desligamento acidental", rai.texto)

    def test_cabecalho_com_numero_de_pagina_e_linha_curta_saem_dos_trechos(self) -> None:
        def pagina(numero: int, corpo: str) -> tuple[int, str]:
            return numero, "\n".join(["O&M", f"Procedimentos de Tempo Real | ITOM AGD 01 | 00 | {numero} de 4", corpo])

        paginas = [
            pagina(1, "1 OBJETIVO\nEstabelecer os procedimentos de operação."),
            pagina(2, "2 DISJUNTORES\nO disjuntor 522 opera fechado."),
            pagina(3, "3 SECCIONADORAS\nNão operar seccionadora em carga."),
            pagina(4, "4 REGISTROS\nRegistrar no RDO."),
        ]
        trechos = fatiar(_documento(paginas, codigo="ITOM AGD 01", titulo="Tempo Real"))
        texto = "\n".join(trecho.texto for trecho in trechos)
        self.assertNotIn("O&M", texto)
        self.assertNotIn("de 4", texto)
        self.assertIn("O disjuntor 522 opera fechado.", texto)
        self.assertEqual([trecho.secao for trecho in trechos], ["1 OBJETIVO", "2 DISJUNTORES", "3 SECCIONADORAS", "4 REGISTROS"])

    def test_layout_marca_colunas_da_tabela(self) -> None:
        bruto = "Intervenção                  7 dias úteis        4 dias úteis\n            Programada – Alta                 15 dias úteis"
        texto = _compactar_layout(bruto)
        self.assertIn("Intervenção | 7 dias úteis | 4 dias úteis", texto)
        self.assertIn("Programada – Alta | 15 dias úteis", texto)

    def test_linha_de_prazo_nao_abre_secao(self) -> None:
        pagina = "\n".join(
            [
                "10.2 VALIDAÇÃO DO RM",
                "Classificação da Intervenção",
                "Intervenção Programada – Baixa complexidade",
                "7 dias úteis 4 dias úteis antes da data",
                "Intervenção Programada – Alta complexidade",
                "15 dias úteis 9 dias úteis antes da data",
            ]
        )
        documento = Documento(
            arquivo="itom.pdf",
            caminho=Path("itom.pdf"),
            texto="",
            paginas=[(16, pagina)],
            codigo="ITOM-01",
            revisao="02",
            vigente=True,
            titulo="ITOM",
        )
        trechos = fatiar(documento)
        secoes = {trecho.secao for trecho in trechos}
        self.assertNotIn("7 dias úteis 4 dias úteis antes da data", secoes)
        self.assertNotIn("15 dias úteis 9 dias úteis antes da data", secoes)
        junto = next(trecho for trecho in trechos if "15 dias úteis" in trecho.texto)
        self.assertIn("Baixa complexidade", junto.texto)
        self.assertIn("Alta complexidade", junto.texto)
        self.assertTrue(junto.secao.startswith("10.2"))

    def test_figura_fica_no_quadro_e_nao_abre_secao(self) -> None:
        documento = Documento(
            arquivo="itom.pdf",
            caminho=Path("itom.pdf"),
            texto="",
            paginas=[
                (
                    11,
                    "\n".join(
                        [
                            "5.1.1 SOLICITAÇÃO DA SI-PS",
                            "O prazo ocorre conforme o quadro a seguir, de acordo com a complexidade.",
                            "<<<FIGURA pagina=11>>>",
                            "Baixa complexidade | 7 dias úteis",
                            "7 AUTORIZAÇÃO QUE NAO DEVE VIRAR SECAO",
                            "<<<FIM FIGURA>>>",
                            "5.2 EMISSÃO DA SI-PS",
                            "O emitente é o responsável.",
                        ]
                    ),
                )
            ],
            codigo="ITOM-01",
            revisao="02",
            vigente=True,
            titulo="ITOM",
        )
        documento = anexar_figuras(documento, {})
        trechos = fatiar(documento)
        quadro = next(trecho for trecho in trechos if trecho.secao.startswith("5.1.1"))
        self.assertIn("Transcrição da figura da página 11", quadro.texto)
        self.assertIn("Fim da transcrição da figura.", quadro.texto)
        self.assertIn("Baixa complexidade | 7 dias úteis", quadro.texto)
        self.assertIn("7 AUTORIZAÇÃO QUE NAO DEVE VIRAR SECAO", quadro.texto)
        emissao = next(trecho for trecho in trechos if trecho.secao.startswith("5.2"))
        self.assertNotIn("Baixa complexidade", emissao.texto)

    def test_quadro_anunciado_no_meio_da_frase_nao_engole_as_secoes_seguintes(self) -> None:
        texto = "\n".join(
            [
                "5.1.1 SOLICITAÇÃO DA SI-PS",
                "O prazo ocorre conforme o quadro a seguir, de acordo com a complexidade.",
                "5.2 EMISSÃO DA SI-PS",
                "O emitente é o responsável.",
                "18. RESPONSABILIDADES E AUTORIDADES",
                "O liberador confere a SI-PS.",
            ]
        )
        documento = Documento(
            arquivo="itom.pdf",
            caminho=Path("itom.pdf"),
            texto=texto,
            paginas=[(11, texto)],
            codigo="ITOM-01",
            revisao="02",
            vigente=True,
            titulo="ITOM",
        )
        documento = anexar_figuras(documento, {11: "Linha: Alta | Coluna: Programada | Valor: 15 dias úteis"})
        self.assertIn("a seguir, de acordo com a complexidade.\n<<<FIGURA", documento.paginas[0][1])
        trechos = fatiar(documento)
        secoes = [trecho.secao for trecho in trechos]
        self.assertEqual(secoes, ["5.1.1 SOLICITAÇÃO DA SI-PS", "5.2 EMISSÃO DA SI-PS", "18 RESPONSABILIDADES E AUTORIDADES"])
        self.assertIn("15 dias úteis", trechos[0].texto)
        self.assertNotIn("15 dias úteis", trechos[1].texto)

    def test_pagina_de_sumario_vira_um_trecho_so(self) -> None:
        sumario = "\n".join(
            ["SUMÁRIO"]
            + [f"{numero}.1 TITULO DO ITEM {numero} {'.' * 40} {numero + 3}" for numero in range(5, 11)]
        )
        corpo = "5.1 TITULO DO ITEM 5\nConteúdo real do item."
        documento = Documento(
            arquivo="itom.pdf",
            caminho=Path("itom.pdf"),
            texto=f"{sumario}\n{corpo}",
            paginas=[(2, sumario), (8, corpo)],
            codigo="ITOM-01",
            revisao="02",
            vigente=True,
            titulo="ITOM",
        )
        trechos = fatiar(documento)
        self.assertEqual([trecho.secao for trecho in trechos], ["Sumário", "5.1 TITULO DO ITEM 5"])
        self.assertEqual(trechos[1].pagina, "8")

    def test_sumario_avisa_secao_que_nao_virou_titulo(self) -> None:
        sumario = "\n".join(
            ["SUMÁRIO"]
            + [f"{numero}. TITULO {numero} {'.' * 40} {numero + 3}" for numero in range(14, 19)]
            + [f"17. | REGISTROS {'.' * 40} 28"]
        )
        corpo = "\n".join(
            [
                "14. TITULO 14",
                "Texto.",
                "15. TITULO 15",
                "Texto.",
                "16. TITULO 16",
                "Texto.",
                "17. | REGISTROS",
                "• | Mensagem Operativa | MO;",
                "Texto sobre a 18 TITULO 18 no meio da frase.",
            ]
        )
        documento = Documento(
            arquivo="itom.pdf",
            caminho=Path("itom.pdf"),
            texto=f"{sumario}\n{corpo}",
            paginas=[(2, sumario), (8, corpo)],
            codigo="ITOM-01",
            revisao="02",
            vigente=True,
            titulo="ITOM",
        )
        trechos = fatiar(documento)
        self.assertIn("17 REGISTROS", [trecho.secao for trecho in trechos])
        self.assertEqual(conferir_sumario(trechos), {"itom.pdf": {"no_sumario": 5, "ausentes": ["18"]}})

    def test_verbete_comecando_por_secao_nao_vira_titulo(self) -> None:
        pagina = "\n".join(
            [
                "Anexo V – GLOSSÁRIO E TERMOS T | ÉCNICOS",
                "Seção de Barramento | Parte de uma subestação.",
                "Seção de Transformador | Parte que compreende um transformador.",
                "Sincronizar | Acoplar eletricamente uma unidade geradora.",
            ]
        )
        trechos = fatiar(_documento([(15, pagina)]))
        self.assertEqual([trecho.secao for trecho in trechos], ["Anexo V GLOSSÁRIO E TERMOS TÉCNICOS"])
        self.assertIn("Sincronizar", trechos[0].texto)

    def test_titulo_numerado_fora_do_sumario_nao_abre_secao(self) -> None:
        sumario = "\n".join(["SUMÁRIO"] + [f"{numero}. TITULO {numero} {'.' * 40} 3" for numero in range(1, 6)])
        corpo = "\n".join(["4. TITULO 4", "Texto do item 4.", "7 AUTORIZAÇÃO DE TESTE", "Continua no item 4."])
        trechos = fatiar(_documento([(2, sumario), (3, corpo)]))
        self.assertEqual([trecho.secao for trecho in trechos], ["Sumário", "4 TITULO 4"])
        self.assertIn("7 AUTORIZAÇÃO DE TESTE", trechos[1].texto)

    def test_trecho_tem_documento_e_secoes_acima(self) -> None:
        pagina = "\n".join(
            [
                "4 PROCEDIMENTOS",
                "4.3 PROCEDIMENTOS GERAIS",
                "4.3.3 MENSAGENS DE MANOBRAS",
                "O receptor deve repetir a mensagem.",
            ]
        )
        documento = _documento([(5, pagina)], codigo="ITOM EGP 05 Rev.00 (2)", titulo="Comunicação na Operação")
        trechos = fatiar(documento)
        self.assertEqual([trecho.secao for trecho in trechos], ["4.3.3 MENSAGENS DE MANOBRAS"])
        linhas = trechos[0].texto.splitlines()
        self.assertEqual(linhas[0], "Documento: ITOM EGP 05 Rev.00 – Comunicação na Operação")
        self.assertEqual(linhas[1], "Seção: 4.3.3 MENSAGENS DE MANOBRAS")
        self.assertEqual(linhas[2], "Dentro de: 4 PROCEDIMENTOS > 4.3 PROCEDIMENTOS GERAIS")
        self.assertEqual(linha_documento(_documento([(1, "")], codigo="IT-1", titulo="IT-1")), "Documento: IT-1")

    def test_janela_corta_no_fim_da_frase(self) -> None:
        verbetes = [f"Termo {numero} | Definição do termo {numero} que ocupa uma linha inteira\ne termina aqui." for numero in range(80)]
        trechos = fatiar(_documento([(15, "Anexo V – GLOSSÁRIO\n" + "\n".join(verbetes))]))
        self.assertGreater(len(trechos), 1)
        for trecho in trechos[:-1]:
            self.assertTrue(trecho.texto.endswith("e termina aqui."), trecho.texto[-60:])

    def test_word_le_tabela_titulo_em_portugues_e_numeracao(self) -> None:
        import docx
        from docx.enum.style import WD_STYLE_TYPE
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn

        def numerar(paragrafo) -> None:
            numeracao = OxmlElement("w:numPr")
            for nome in ("w:ilvl", "w:numId"):
                elemento = OxmlElement(nome)
                elemento.set(qn("w:val"), "1")
                numeracao.append(elemento)
            paragrafo._p.get_or_add_pPr().append(numeracao)

        arquivo = docx.Document()
        arquivo.add_paragraph("Objetivo 3", style=arquivo.styles.add_style("TOC 1", WD_STYLE_TYPE.PARAGRAPH))
        numerar(arquivo.add_heading("Objetivo", level=1))
        arquivo.add_paragraph("Texto do objetivo.")
        numerar(arquivo.add_heading("Procedimentos", level=1))
        numerar(arquivo.add_paragraph("Mensagens de manobras", style=arquivo.styles.add_style("Título 2", WD_STYLE_TYPE.PARAGRAPH)))
        arquivo.add_paragraph("Repetir a mensagem.")
        tabela = arquivo.add_table(rows=1, cols=2)
        tabela.cell(0, 0).text = "Sincronizar"
        tabela.cell(0, 1).text = "Acoplar eletricamente."
        with tempfile.TemporaryDirectory() as temporario:
            caminho = Path(temporario) / "it-word.docx"
            arquivo.save(caminho)
            documento = carregar_arquivo(caminho)
        trechos = fatiar(documento)
        self.assertEqual([trecho.secao for trecho in trechos], ["1 Objetivo", "2.1 Mensagens de manobras"])
        self.assertIn("Sincronizar | Acoplar eletricamente.", trechos[1].texto)
        self.assertIn("Dentro de: 2 Procedimentos", trechos[1].texto)
        self.assertNotIn("Objetivo 3", documento.texto)
        self.assertEqual(documento.titulo, "it-word")

    def test_transcricao_guardada_nao_chama_o_gpt(self) -> None:
        figura = Figura("itom.pdf", 11, "nada-gravado", b"png", "image/png")
        with tempfile.TemporaryDirectory() as temporario:
            self.assertEqual(transcricoes_guardadas(Path(temporario), [figura]), {})
            self.assertFalse((Path(temporario) / "visao").exists())

    def test_marcador_de_fim_com_texto_colado_fecha_a_figura(self) -> None:
        texto = "5.1.1 SOLICITAÇÃO\n<<<FIGURA pagina=11>>>\nLinha: A | Valor: 1\n<<<FIM FIGURA>>>, resto da frase\n5.2 EMISSÃO DA SI-PS\nTexto."
        documento = Documento(
            arquivo="itom.pdf",
            caminho=Path("itom.pdf"),
            texto=texto,
            paginas=[(11, texto)],
            codigo="ITOM-01",
            revisao="02",
            vigente=True,
            titulo="ITOM",
        )
        trechos = fatiar(documento)
        self.assertEqual([trecho.secao for trecho in trechos], ["5.1.1 SOLICITAÇÃO", "5.2 EMISSÃO DA SI-PS"])
        self.assertIn(", resto da frase", trechos[0].texto)

    def test_logotipo_repetido_fica_de_fora(self) -> None:
        pdf = next(RAIZ.joinpath("acervo").glob("ITOM EGP*.pdf"), None)
        if pdf is None:
            self.skipTest("ITOM EGP ausente")
        figuras = figuras_de_conteudo(pdf)
        paginas = {figura.pagina for figura in figuras}
        self.assertIn(11, paginas)
        self.assertLess(len(figuras), 20)
        self.assertEqual(len(figuras), len({figura.hash for figura in figuras}))

    def test_transcricao_gravada_nao_chama_de_novo(self) -> None:
        from tempfile import TemporaryDirectory

        figura = Figura("itom.pdf", 11, "abc123", b"png", "image/png")
        chamadas: list[str] = []

        def gerador(item: Figura) -> str:
            chamadas.append(item.hash)
            return "SEM_CONTEUDO"

        with TemporaryDirectory() as temporario:
            data = Path(temporario)
            self.assertEqual(transcrever_figuras(data, [figura], gerador=gerador), {})
            self.assertEqual(transcrever_figuras(data, [figura], gerador=gerador), {})
        self.assertEqual(chamadas, ["abc123"])

    def test_glossario_leva_wtg_ao_anexo(self) -> None:
        self.assertIn("aerogerador", expandir("manutenção de WTG com seccionamento de SU"))
        self.assertIn("desligamento", expandir("seccionamento"))
        self.assertIn("urgência", expandir("de urgências"))
        self.assertEqual(expandir("aerogerador"), "aerogerador")
        anexo = Trecho(
            "1",
            "Anexo 01. CD (Com Desligamento) para aerogerador e seccionadora.",
            "itom.pdf",
            "ITOM-01",
            "02",
            True,
            "ITOM",
            "Anexo 01 SI-PS",
            "35",
        )
        outro = Trecho(
            "2",
            "Linha energizada apenas no barramento.",
            "itom.pdf",
            "ITOM-01",
            "02",
            True,
            "ITOM",
            "11.3.2",
            "20",
        )
        neutro = Trecho(
            "3",
            "Calendário administrativo do escritório central.",
            "itom.pdf",
            "ITOM-01",
            "02",
            True,
            "ITOM",
            "Documento",
            "1",
        )
        indice = Indice("teste", "x", [outro, neutro, anexo], np.zeros((3, 4), dtype=np.float32))
        acertos = buscar(indice, "manutenção de WTG com seccionamento de SU", np.zeros(4, dtype=np.float32), top_k=3)
        self.assertEqual(acertos[0].trecho.secao, "Anexo 01 SI-PS")
        self.assertGreater(acertos[0].bm25, 0)

    def test_continuacao_acompanha_a_parte_seguinte(self) -> None:
        def trecho(secao: str) -> Trecho:
            return Trecho("1", secao, "a.pdf", "NO-01", "00", True, "t", secao, "1")

        trechos = [
            trecho("2.4 RAI (parte 1)"),
            trecho("2.4 RAI (parte 2)"),
            trecho("3 Relatório"),
        ]
        self.assertEqual(_com_continuacao(trechos, [0], limite=10), [0, 1])

    def test_figura_entra_mesmo_quando_o_acerto_e_parte_posterior(self) -> None:
        def trecho(secao: str, texto: str) -> Trecho:
            return Trecho("1", texto, "a.pdf", "ITOM", "02", True, "t", secao, "11")

        trechos = [
            trecho(
                "5.1.1 SOLICITAÇÃO (parte 1)",
                "Transcrição da figura da página 11:\nLinha: Alta | Coluna: Programada | Valor: 15 dias úteis\nFim da transcrição da figura.",
            ),
            trecho("5.1.1 SOLICITAÇÃO (parte 2)", "texto da parte 2"),
            trecho("5.1.1 SOLICITAÇÃO (parte 3)", "texto da parte 3"),
        ]
        self.assertEqual(_com_continuacao(trechos, [2], limite=2), [2, 0])
        self.assertEqual(_com_continuacao(trechos, [2, 1], limite=2), [2, 1, 0])


class CacheVetoresTeste(unittest.TestCase):
    def test_reindexar_so_calcula_textos_novos(self) -> None:
        pedidos: list[list[str]] = []

        def calcular(textos: list[str]) -> list[np.ndarray]:
            pedidos.append(list(textos))
            return [np.full(3, float(len(texto)), dtype=np.float32) for texto in textos]

        with tempfile.TemporaryDirectory() as pasta:
            data_dir = Path(pasta)
            primeiro = _com_cache(data_dir, "modelo", ["a", "bb"], calcular)
            segundo = _com_cache(data_dir, "modelo", ["bb", "ccc", "a"], calcular)
            outro_modelo = _com_cache(data_dir, "outro", ["a"], calcular)

        self.assertEqual(pedidos, [["a", "bb"], ["ccc"], ["a"]])
        self.assertEqual([float(vetor[0]) for vetor in primeiro], [1.0, 2.0])
        self.assertEqual([float(vetor[0]) for vetor in segundo], [2.0, 3.0, 1.0])
        self.assertEqual(float(outro_modelo[0][0]), 1.0)


if __name__ == "__main__":
    unittest.main()
